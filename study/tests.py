from datetime import timedelta
from io import BytesIO
from types import SimpleNamespace
from tempfile import TemporaryDirectory
from unittest.mock import patch
import httpx
from openai import APIConnectionError, APIStatusError, APITimeoutError, RateLimitError
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document
from .demo import create_demo
from .forms import CourseForm, LectureForm
from .models import Lecture, Question, StudySession
from .services import LectureSummary, QuestionBank, generate_questions, generate_summary, question_selection, start_session, streak, submit_session

class StudyTests(TestCase):
    def setUp(self):
        media = TemporaryDirectory()
        self.addCleanup(media.cleanup)
        media_settings = override_settings(MEDIA_ROOT=media.name)
        media_settings.enable()
        self.addCleanup(media_settings.disable)
        self.user = User.objects.create_user('student', password='a-strong-test-password')
        self.other = User.objects.create_user('other', password='another-test-password')
        self.course = create_demo(self.user)
        self.client.force_login(self.user)

    def answers(self, session):
        return {f'question_{i.pk}': str(i.question.correct_index) for i in session.items.select_related('question')}

    def test_page_rendering(self):
        for url in ['/', reverse('course_detail', args=[self.course.pk]), reverse('course_edit', args=[self.course.pk]),
                    reverse('upload', args=[self.course.pk]), reverse('course_create')]:
            self.assertEqual(self.client.get(url).status_code, 200)
        self.client.logout()
        self.assertEqual(self.client.get('/').status_code, 302)
        self.assertEqual(self.client.get(reverse('login')).status_code, 200)
        self.assertEqual(self.client.get(reverse('signup')).status_code, 200)

    def test_signup_and_course_creation(self):
        self.client.logout()
        response = self.client.post(reverse('signup'), {'username': 'newstudent', 'password1': 'Study-Account-284!',
                                                       'password2': 'Study-Account-284!'})
        self.assertEqual(response.status_code, 302)
        response = self.client.post(reverse('course_create'), {'name': 'Algorithms', 'daily_minutes': 7,
            'exam_date': timezone.localdate() + timedelta(days=20)})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(User.objects.get(username='newstudent').courses.filter(name='Algorithms').exists())

    def test_owner_isolation(self):
        session = start_session(self.course, 'daily')
        self.client.force_login(self.other)
        for name, args in [('course_detail', [self.course.pk]), ('course_edit', [self.course.pk]),
                           ('upload', [self.course.pk]), ('quiz', [session.pk])]:
            self.assertEqual(self.client.get(reverse(name, args=args)).status_code, 404)
        for name, args in [('start', [self.course.pk, 'daily']), ('generate', [self.course.lectures.first().pk]),
                           ('quiz', [session.pk])]:
            self.assertEqual(self.client.post(reverse(name, args=args)).status_code, 404)

    def test_daily_budget_and_deadline(self):
        self.course.daily_minutes = 3
        self.assertEqual(len(question_selection(self.course, 'daily')), 3)
        self.course.daily_minutes = 10
        self.course.exam_date = timezone.localdate()
        self.assertEqual(len(question_selection(self.course, 'daily')), 10)
        self.course.exam_date = timezone.localdate() + timedelta(days=100)
        self.assertEqual(len(question_selection(self.course, 'daily')), 5)

    def test_weak_due_questions_first(self):
        q = Question.objects.filter(topic__course=self.course).last()
        Question.objects.filter(topic__course=self.course).update(mastery=90)
        q.mastery = 0
        q.save()
        self.assertEqual(question_selection(self.course, 'daily')[0].pk, q.pk)

    def test_weekly_coverage_and_period(self):
        weekly = start_session(self.course, 'weekly')
        self.assertEqual(weekly.items.count(), 12)
        self.assertEqual(len(set(weekly.items.values_list('question__topic_id', flat=True))), 3)
        self.assertEqual(weekly.period_date.weekday(), 0)
        self.assertEqual(start_session(self.course, 'weekly').pk, weekly.pk)

    def test_quiz_submission_and_replay(self):
        session = start_session(self.course, 'daily')
        self.assertEqual(start_session(self.course, 'daily').pk, session.pk)
        url = reverse('quiz', args=[session.pk])
        self.assertEqual(self.client.get(url).status_code, 200)
        self.assertNotContains(self.client.get(url), 'Why this answer?')
        self.assertEqual(self.client.post(url, self.answers(session)).status_code, 302)
        session.refresh_from_db()
        self.assertEqual(session.score, 100)
        self.assertIsNotNone(session.completed_at)
        q = session.items.first().question
        self.assertEqual(q.mastery, 25)
        self.assertGreater(q.due_date, timezone.localdate())
        submit_session(session, self.answers(session))
        q.refresh_from_db()
        self.assertEqual(q.attempts, 1)
        self.assertContains(self.client.get(url), 'Why this answer?')
        self.assertEqual(streak(self.user), 1)

    def test_invalid_submission_is_atomic(self):
        session = start_session(self.course, 'daily')
        data = self.answers(session)
        data[f'question_{session.items.last().pk}'] = '999'
        with self.assertRaises(ValueError):
            submit_session(session, data)
        self.assertFalse(session.items.filter(selected_index__isnull=False).exists())
        self.assertFalse(Question.objects.filter(attempts__gt=0).exists())
        with self.assertRaises(ValueError):
            submit_session(session, {})

    def test_wrong_answer_reduces_mastery(self):
        session = start_session(self.course, 'daily')
        item = session.items.first()
        Question.objects.filter(pk=item.question_id).update(mastery=75)
        data = self.answers(session)
        data[f'question_{item.pk}'] = str((item.question.correct_index + 1) % 4)
        submit_session(session, data)
        q = Question.objects.get(pk=item.question_id)
        self.assertEqual(q.mastery, 40)
        self.assertEqual(q.due_date, timezone.localdate() + timedelta(days=1))

    def test_streak_counts_days_not_sessions_and_expires(self):
        now = timezone.now()
        for days in [1, 1, 2, 4]:
            StudySession.objects.create(course=self.course, kind=f'x{StudySession.objects.count()}',
                period_date=timezone.localdate(), completed_at=now - timedelta(days=days))
        self.assertEqual(streak(self.user), 2)
        StudySession.objects.all().update(completed_at=now - timedelta(days=3))
        self.assertEqual(streak(self.user), 0)

    @patch('study.views.generate_summary')
    def test_upload_docx_and_reject_bad_file(self, summary):
        stream = BytesIO()
        doc = Document()
        doc.add_paragraph('Primary keys uniquely identify rows. Foreign keys reference related tables. ' * 3)
        doc.save(stream)
        response = self.client.post(reverse('upload', args=[self.course.pk]), {
            'title': 'New lecture', 'file': SimpleUploadedFile('lecture.docx', stream.getvalue())})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(self.course.lectures.filter(title='New lecture').exists())
        lecture = self.course.lectures.get(title='New lecture')
        summary.assert_called_once()
        self.assertEqual(summary.call_args.args[0].pk, lecture.pk)
        with lecture.file.open('rb') as saved:
            self.assertEqual(saved.read(), stream.getvalue())
        original = self.client.get(reverse('lecture_open', args=[lecture.pk]))
        self.assertEqual(original['Content-Type'], 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        self.assertTrue(original['Content-Disposition'].startswith('attachment;'))
        self.assertEqual(b''.join(original.streaming_content), stream.getvalue())
        download = self.client.get(reverse('lecture_open', args=[lecture.pk]) + '?download=1')
        self.assertEqual(b''.join(download.streaming_content), stream.getvalue())
        for name, content in [('bad.exe', b'not a document'), ('broken.pdf', b'broken')]:
            form = LectureForm({'title': 'Bad'}, {'file': SimpleUploadedFile(name, content)})
            self.assertFalse(form.is_valid())

    @override_settings(OPENAI_API_KEY='test-key')
    @patch('study.services.OpenAI')
    def test_summary_saved_rendered_escaped_and_reused(self, client):
        lecture = self.course.lectures.first()
        summary = LectureSummary(overview='Lecture overview <script>alert(1)</script>',
            concepts=[{'name': 'Key', 'explanation': 'A unique row identifier.'}],
            key_takeaways=['Distinguish primary keys from foreign keys.'])
        client.return_value.__enter__.return_value.responses.parse.return_value = SimpleNamespace(output_parsed=summary)
        response = self.client.post(reverse('summary_generate', args=[lecture.pk]), follow=True)
        self.assertContains(response, 'A unique row identifier.')
        self.assertContains(response, '&lt;script&gt;')
        self.assertNotContains(response, '<script>alert(1)</script>')
        lecture.refresh_from_db()
        self.assertEqual(lecture.summary, {**summary.model_dump(), 'language': 'en'})
        generate_summary(lecture)
        client.assert_called_once()
        self.assertTrue(lecture.generated)  # Demo question bank is preserved.
        kwargs = client.return_value.__enter__.return_value.responses.parse.call_args.kwargs
        self.assertIs(kwargs['text_format'], LectureSummary)
        self.assertIn('entire study guide in English', kwargs['input'][0]['content'])
        self.assertEqual(kwargs['input'][1]['content'], lecture.text)

    @override_settings(OPENAI_API_KEY='test-key')
    @patch('study.services.OpenAI')
    def test_summary_refusal_and_connection_error_can_be_retried(self, client):
        lecture = self.course.lectures.first()
        parse = client.return_value.__enter__.return_value.responses.parse
        parse.return_value = SimpleNamespace(output_parsed=None)
        response = self.client.post(reverse('summary_generate', args=[lecture.pk]), follow=True)
        self.assertContains(response, 'OpenAI returned no summary')
        parse.side_effect = APIConnectionError(request=httpx.Request('POST', 'https://api.openai.com/v1/responses'))
        response = self.client.post(reverse('summary_generate', args=[lecture.pk]), follow=True)
        self.assertContains(response, 'Could not reach OpenAI')
        lecture.refresh_from_db()
        self.assertEqual(lecture.summary, {})
        self.assertContains(response, 'Create summary')

    @override_settings(OPENAI_API_KEY='test-key')
    @patch('study.services.OpenAI')
    def test_legacy_summary_replaced_in_english_only_on_success(self, client):
        lecture = self.course.lectures.first()
        old_summary = {'overview': 'Старый конспект', 'concepts': [], 'key_takeaways': []}
        lecture.summary = old_summary
        lecture.save(update_fields=['summary'])
        url = reverse('lecture_summary', args=[lecture.pk])
        response = self.client.get(url)
        self.assertContains(response, 'Create summary')
        self.assertNotContains(response, 'English summary')
        self.assertNotContains(response, 'Старый конспект')
        client.assert_not_called()
        parse = client.return_value.__enter__.return_value.responses.parse
        parse.return_value = SimpleNamespace(output_parsed=None)
        self.client.post(reverse('summary_generate', args=[lecture.pk]))
        lecture.refresh_from_db()
        self.assertEqual(lecture.summary, old_summary)
        parse.return_value = SimpleNamespace(output_parsed=LectureSummary(
            overview='English overview', concepts=[{'name': 'Key', 'explanation': 'Unique identifier.'}],
            key_takeaways=['Understand keys.']))
        response = self.client.post(reverse('summary_generate', args=[lecture.pk]), follow=True)
        self.assertContains(response, 'English overview')
        lecture.refresh_from_db()
        self.assertEqual(lecture.summary['language'], 'en')

    @patch('study.views.generate_summary', side_effect=ValueError('Summary unavailable'))
    def test_failed_summary_does_not_lose_upload(self, summary):
        stream = BytesIO()
        doc = Document()
        doc.add_paragraph('Lecture on primary and foreign keys. ' * 10)
        doc.save(stream)
        response = self.client.post(reverse('upload', args=[self.course.pk]), {
            'title': 'Saved despite API error', 'file': SimpleUploadedFile('lecture.docx', stream.getvalue())}, follow=True)
        self.assertContains(response, 'Lecture saved')
        lecture = self.course.lectures.get(title='Saved despite API error')
        self.assertTrue(lecture.file.storage.exists(lecture.file.name))
        self.assertEqual(lecture.summary, {})

    def test_lecture_access_and_legacy_text_file(self):
        lecture = self.course.lectures.first()
        response = self.client.get(reverse('lecture_open', args=[lecture.pk]))
        self.assertEqual(response['Content-Type'], 'text/plain; charset=utf-8')
        self.assertIn('.txt', response['Content-Disposition'])
        self.assertTrue(response['Content-Disposition'].startswith('attachment;'))
        self.assertEqual(b''.join(response.streaming_content).decode('utf-8'), lecture.text)
        page = self.client.get(reverse('course_detail', args=[self.course.pk]))
        self.assertContains(page, reverse('lecture_open', args=[lecture.pk]))
        self.assertContains(page, 'download>')
        self.assertNotContains(page, 'target="_blank"')
        self.client.force_login(self.other)
        for name in ['lecture_open', 'lecture_summary']:
            self.assertEqual(self.client.get(reverse(name, args=[lecture.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse('summary_generate', args=[lecture.pk])).status_code, 404)
        self.client.logout()
        self.assertEqual(self.client.get(reverse('lecture_open', args=[lecture.pk])).status_code, 302)

    def test_pdf_download_and_missing_file(self):
        lecture = self.course.lectures.first()
        lecture.original_name = 'lecture.pdf'
        content = b'%PDF-1.4\nTest document'
        lecture.file.save('lecture.pdf', SimpleUploadedFile('lecture.pdf', content))
        url = reverse('lecture_open', args=[lecture.pk])
        response = self.client.get(url)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response['Content-Disposition'].startswith('attachment;'))
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        self.assertEqual(b''.join(response.streaming_content), content)
        lecture.file.storage.delete(lecture.file.name)
        self.assertEqual(self.client.get(url).status_code, 404)

    def test_delete_lecture_removes_file_questions_and_unused_topics(self):
        lecture = self.course.lectures.first()
        lecture.file.save('lecture.pdf', SimpleUploadedFile('lecture.pdf', b'%PDF-1.4'))
        storage, filename = lecture.file.storage, lecture.file.name
        question_ids = list(lecture.questions.values_list('pk', flat=True))
        topic_ids = list(lecture.questions.values_list('topic_id', flat=True))
        # A shared topic and its other lecture must survive.
        other_lecture = Lecture.objects.create(course=self.course, title='Other', text='Text', original_name='other.pdf')
        question = lecture.questions.first()
        question.pk = None
        question.lecture = other_lecture
        question.save()
        start_session(self.course, 'daily')
        url = reverse('lecture_delete', args=[lecture.pk])
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(url, follow=True)
        self.assertContains(response, 'Lecture deleted.')
        self.assertFalse(Lecture.objects.filter(pk=lecture.pk).exists())
        self.assertFalse(Question.objects.filter(pk__in=question_ids).exists())
        self.assertFalse(storage.exists(filename))
        self.assertTrue(other_lecture.questions.filter(pk=question.pk).exists())
        self.assertEqual(list(self.course.topics.filter(pk__in=topic_ids).values_list('pk', flat=True)), [question.topic_id])

    def test_delete_requires_owner_post_and_csrf(self):
        from django.test import Client
        lecture = self.course.lectures.first()
        url = reverse('lecture_delete', args=[lecture.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        self.client.force_login(self.other)
        self.assertEqual(self.client.post(url).status_code, 404)
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        self.assertEqual(client.post(url).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.post(url).status_code, 302)
        self.assertTrue(Lecture.objects.filter(pk=lecture.pk).exists())

    def test_course_validation(self):
        for date, minutes in [(timezone.localdate()-timedelta(days=1), 10), (timezone.localdate(), 11),
                              (timezone.localdate(), 0)]:
            self.assertFalse(CourseForm({'name': 'Course', 'exam_date': date, 'daily_minutes': minutes}).is_valid())

    @override_settings(OPENAI_API_KEY='')
    def test_missing_key(self):
        lecture = Lecture.objects.create(course=self.course, title='Test', text='Lecture content', original_name='test.pdf')
        with self.assertRaisesMessage(ValueError, 'OPENAI_API_KEY'):
            generate_questions(lecture)
        self.assertFalse(lecture.questions.exists())

    @override_settings(OPENAI_API_KEY='test-key', OPENAI_MODEL='gpt-6-luna')
    @patch('study.services.OpenAI')
    def test_generated_bank_saved_once(self, client):
        lecture = Lecture.objects.create(course=self.course, title='Test', text='Lecture content', original_name='test.pdf')
        bank = QuestionBank(questions=[{'topic': 'Keys', 'prompt': 'What is unique?', 'choices': ['Primary key', 'Name', 'Age', 'City'],
                                       'correct_index': 0, 'explanation': 'Primary keys identify unique rows.'}])
        client.return_value.__enter__.return_value.responses.parse.return_value = SimpleNamespace(output_parsed=bank)
        self.assertEqual(generate_questions(lecture), 1)
        lecture.refresh_from_db()
        self.assertEqual(generate_questions(lecture), 0)
        self.assertEqual(lecture.questions.count(), 1)
        call = client.return_value.__enter__.return_value.responses.parse.call_args.kwargs
        self.assertEqual(call['model'], 'gpt-6-luna')
        self.assertEqual(call['reasoning'], {'effort': 'none'})
        self.assertEqual(call['input'][1]['content'], lecture.text)
        self.assertIn('untrusted', call['input'][0]['content'])
        self.assertIs(call['text_format'], QuestionBank)
        client.return_value.__exit__.assert_called_once()

    @override_settings(OPENAI_API_KEY='test-key')
    @patch('study.services.OpenAI')
    def test_refusal_keeps_lecture_retryable(self, client):
        lecture = Lecture.objects.create(course=self.course, title='Test', text='Lecture content', original_name='test.pdf')
        client.return_value.__enter__.return_value.responses.parse.return_value = SimpleNamespace(output_parsed=None)
        with self.assertRaises(ValueError):
            generate_questions(lecture)
        lecture.refresh_from_db()
        self.assertFalse(lecture.generated)

    def test_mutations_require_post(self):
        for name, args in [('start', [self.course.pk, 'daily']), ('generate', [self.course.lectures.first().pk]),
                           ('summary_generate', [self.course.lectures.first().pk]), ('demo', [])]:
            self.assertEqual(self.client.get(reverse(name, args=args)).status_code, 405)

    def test_demo_is_repeatable(self):
        self.assertEqual(create_demo(self.user).pk, self.course.pk)
        self.assertEqual(Question.objects.count(), 12)

    @override_settings(OPENAI_API_KEY='test-key')
    @patch('study.services.OpenAI')
    def test_openai_duplicate_choices_rejected(self, client):
        lecture = Lecture.objects.create(course=self.course, title='Test', text='Lecture', original_name='test.pdf')
        bank = QuestionBank(questions=[{'topic': 'Keys', 'prompt': 'What is unique?', 'choices': ['Key', 'key ', 'Age', 'City'],
                                       'correct_index': 0, 'explanation': 'Explanation.'}])
        client.return_value.__enter__.return_value.responses.parse.return_value = SimpleNamespace(output_parsed=bank)
        with self.assertRaisesMessage(ValueError, 'invalid'):
            generate_questions(lecture)
        self.assertFalse(lecture.questions.exists())

    @override_settings(OPENAI_API_KEY='test-key')
    @patch('study.services.OpenAI')
    def test_openai_provider_and_network_errors_are_recoverable(self, client):
        lecture = Lecture.objects.create(course=self.course, title='Test', text='Lecture', original_name='test.pdf')
        request = httpx.Request('POST', 'https://api.openai.com/v1/responses')
        cases = [
            (RateLimitError('private provider message', response=httpx.Response(429, request=request), body={}), 'quota'),
            (APIStatusError('private provider message', response=httpx.Response(403, request=request), body={}), 'model access'),
            (APIConnectionError(message='private connection details', request=request), 'internet connection'),
        ]
        for error, expected in cases:
            client.return_value.__enter__.return_value.responses.parse.side_effect = error
            response = self.client.post(reverse('generate', args=[lecture.pk]), follow=True)
            self.assertContains(response, expected)
            self.assertNotContains(response, 'private provider message')
            self.assertNotContains(response, 'private connection details')
            lecture.refresh_from_db()
            self.assertFalse(lecture.generated)
            self.assertFalse(lecture.questions.exists())

    @override_settings(OPENAI_API_KEY='test-key')
    @patch('study.services.OpenAI')
    def test_openai_server_error_has_safe_message(self, client):
        lecture = Lecture.objects.create(course=self.course, title='Test', text='Lecture', original_name='test.pdf')
        request = httpx.Request('POST', 'https://api.openai.com/v1/responses')
        client.return_value.__enter__.return_value.responses.parse.side_effect = APIStatusError(
            'private provider message', response=httpx.Response(503, request=request), body={})
        response = self.client.post(reverse('generate', args=[lecture.pk]), follow=True)
        self.assertContains(response, 'temporarily unavailable')
        self.assertNotContains(response, 'private provider message')

    @override_settings(OPENAI_API_KEY='test-key')
    @patch('study.services.OpenAI')
    def test_openai_timeout_has_specific_message(self, client):
        lecture = Lecture.objects.create(course=self.course, title='Test', text='Lecture', original_name='test.pdf')
        request = httpx.Request('POST', 'https://api.openai.com/v1/responses')
        client.return_value.__enter__.return_value.responses.parse.side_effect = APITimeoutError(request=request)
        response = self.client.post(reverse('generate', args=[lecture.pk]), follow=True)
        self.assertContains(response, 'took too long')
        lecture.refresh_from_db()
        self.assertFalse(lecture.generated)
